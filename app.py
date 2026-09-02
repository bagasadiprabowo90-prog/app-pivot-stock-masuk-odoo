"""Aplikasi Streamlit untuk pivot file Transfer (stock.picking) dari Odoo."""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO
import re
from typing import Any

import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


REQUIRED_HEADERS = {
    "Operations/Product": "product",
    "Operations/Lot/Serial Number": "batch",
    "Operations/Qty Done": "qty",
}
MAX_FILE_SIZE = 10 * 1024 * 1024


def displayed_value(cell: Any) -> str:
    """Return the visible identifier while retaining Excel zero-padding such as 000123."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()

    if isinstance(value, (int, float)) and float(value).is_integer():
        number_format = str(cell.number_format or "").strip()
        if re.fullmatch(r"0+", number_format):
            return f"{int(value):0{len(number_format)}d}"
        return str(int(value))

    return str(value).strip()


def parse_quantity(value: Any) -> float | None:
    """Parse native or Indonesian-formatted quantities without changing identifiers."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(" ", "")
    if not text:
        return None

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return None


def read_and_pivot(uploaded_file: Any) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(BytesIO(uploaded_file.getvalue()), read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            raise ValueError("File Excel kosong atau rusak.")

        sheet = workbook.worksheets[0]
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = [displayed_value(cell) for cell in header_cells]
        missing_headers = [name for name in REQUIRED_HEADERS if name not in headers]
        if missing_headers:
            raise ValueError(
                "Kolom wajib tidak ditemukan: " + ", ".join(missing_headers) + "."
            )

        column_indexes = {
            key: headers.index(header_name)
            for header_name, key in REQUIRED_HEADERS.items()
        }
        grouped: dict[tuple[str, str], float] = defaultdict(float)
        source_line_count = 0
        skipped_row_count = 0
        cleaned_prefix_count = 0
        source_total_qty = 0.0
        total_source_rows = 0

        for row in sheet.iter_rows(min_row=2):
            if not any(cell.value is not None for cell in row):
                continue
            total_source_rows += 1

            product = displayed_value(row[column_indexes["product"]])
            raw_batch = displayed_value(row[column_indexes["batch"]])
            quantity = parse_quantity(row[column_indexes["qty"]].value)
            sku_match = re.match(r"^\s*\[([^\]]+)\]", product)

            if not sku_match or quantity is None:
                skipped_row_count += 1
                continue

            batch = re.sub(r"^(?:PP-|P-)", "", raw_batch, flags=re.IGNORECASE).strip()
            if raw_batch and batch != raw_batch:
                cleaned_prefix_count += 1

            sku = sku_match.group(1).strip()
            grouped[(sku, batch)] += quantity
            source_line_count += 1
            source_total_qty += quantity

        pivot_rows = [
            {"SKU Batch": f"{sku}{batch}", "Qty": qty, "SKU": sku, "Batch": batch}
            for (sku, batch), qty in grouped.items()
        ]
        pivot_rows.sort(key=lambda row: (natural_key(row["SKU"]), natural_key(row["Batch"])))

        pivot_total_qty = sum(row["Qty"] for row in pivot_rows)
        if abs(source_total_qty - pivot_total_qty) > 1e-9:
            raise ValueError(
                f"Total Qty tidak cocok: sumber {source_total_qty}, pivot {pivot_total_qty}."
            )

        stats = {
            "total_source_rows": total_source_rows,
            "grouped_rows": source_line_count,
            "skipped_rows": skipped_row_count,
            "cleaned_prefix_count": cleaned_prefix_count,
            "pivot_row_count": len(pivot_rows),
            "pivot_total_qty": pivot_total_qty,
        }
        return pivot_rows, stats
    finally:
        workbook.close()


def natural_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def make_output_workbook(pivot_rows: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pivot SKU Batch"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border_color = "0B5D56"
    thin_border = Border(
        top=Side(style="thin", color=border_color),
        left=Side(style="thin", color=border_color),
        right=Side(style="thin", color=border_color),
        bottom=Side(style="thin", color=border_color),
    )
    hair_border = Border(bottom=Side(style="hair", color="E5E7EB"))
    alternate_fill = PatternFill("solid", fgColor="F8FAFC")

    sheet.append(["SKU Batch", "Qty"])
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    sheet.row_dimensions[1].height = 28

    for row_number, item in enumerate(pivot_rows, start=2):
        sheet.append([item["SKU Batch"], item["Qty"]])
        sku_batch_cell = sheet.cell(row=row_number, column=1)
        qty_cell = sheet.cell(row=row_number, column=2)
        sku_batch_cell.number_format = "@"
        qty_cell.number_format = "#,##0.###"
        sku_batch_cell.font = Font(name="Calibri", size=10, color="1F2937")
        qty_cell.font = Font(name="Calibri", size=10, color="1F2937")
        sku_batch_cell.alignment = Alignment(vertical="center")
        qty_cell.alignment = Alignment(vertical="center", horizontal="right")
        sku_batch_cell.border = hair_border
        qty_cell.border = hair_border
        if row_number % 2 == 1:
            sku_batch_cell.fill = alternate_fill
            qty_cell.fill = alternate_fill
        sheet.row_dimensions[row_number].height = 20

    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 16
    sheet.auto_filter.ref = f"A1:B{max(1, len(pivot_rows) + 1)}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def format_quantity(value: float) -> str:
    if float(value).is_integer():
        return f"{int(value):,}".replace(",", ".")
    return f"{value:,.3f}".rstrip("0").rstrip(".").replace(",", "X").replace(".", ",").replace("X", ".")


st.set_page_config(page_title="Pivot Stock Masuk Odoo", page_icon="📊", layout="wide")
st.title("Pivot Stock Masuk Odoo")
st.caption(
    "Upload file Transfer (stock.picking), lalu dapatkan file Excel dengan kolom SKU Batch dan Qty. "
    "Prefix batch P- dan PP- dibersihkan."
)

uploaded_file = st.file_uploader(
    "Pilih file Excel (.xlsx)", type=["xlsx"], help="Maksimal ukuran file 10 MB."
)

if uploaded_file is not None:
    if uploaded_file.size > MAX_FILE_SIZE:
        st.error("File terlalu besar. Maksimal ukuran file adalah 10 MB.")
    else:
        st.info(f"File dipilih: {uploaded_file.name} ({uploaded_file.size / 1024 / 1024:.2f} MB)")
        if st.button("Proses File", type="primary", use_container_width=True):
            try:
                with st.spinner("Memproses file Excel..."):
                    pivot_rows, stats = read_and_pivot(uploaded_file)
                    output_bytes = make_output_workbook(pivot_rows)

                st.success("File berhasil diproses.")
                metrics = st.columns(6)
                metric_values = [
                    ("Baris sumber", stats["total_source_rows"]),
                    ("Diproses", stats["grouped_rows"]),
                    ("Skip", stats["skipped_rows"]),
                    ("Prefix dibersihkan", stats["cleaned_prefix_count"]),
                    ("Baris pivot", stats["pivot_row_count"]),
                    ("Total Qty", format_quantity(stats["pivot_total_qty"])),
                ]
                for column, (label, value) in zip(metrics, metric_values):
                    column.metric(label, value)

                preview_rows = [{"SKU Batch": row["SKU Batch"], "Qty": row["Qty"]} for row in pivot_rows[:100]]
                st.subheader("Preview hasil")
                st.dataframe(preview_rows, use_container_width=True, hide_index=True)
                if len(pivot_rows) > 100:
                    st.caption(f"Menampilkan 100 dari {len(pivot_rows):,} baris pivot.")

                filename = f"Pivot Gabungan SKU Batch dan Qty - {date.today().isoformat()}.xlsx"
                st.download_button(
                    "Download File Pivot",
                    data=output_bytes,
                    file_name=filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True,
                )
            except Exception as error:
                st.error(f"Gagal memproses file: {error}")

st.divider()
st.caption("Aplikasi Streamlit · File diproses sementara dan tidak disimpan setelah sesi berakhir.")
